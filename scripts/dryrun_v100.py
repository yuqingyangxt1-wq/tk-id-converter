"""ID v1.0.0 dry-run：用真实 EasyBoss ID 源表验证转化流程。"""
from __future__ import annotations
import sys, shutil
from pathlib import Path

sys.path.insert(0, "/Users/apple/WorkBuddy/2026-09-08-17-31-48/tk-id-converter-repo")

from app.source_reader import read_source
from app.converter import convert_source
from app.config import default_config

# 真实源表
src = Path("/Users/apple/Desktop/导出#SKU_2026_09_11_17_15_46.xlsx")
out_dir = Path("/tmp/id_dryrun_out")
if out_dir.exists():
    shutil.rmtree(out_dir)

settings = default_config()["product_xlsx_settings"]
print("--- ID DEFAULT SETTINGS ---")
for k in ["title_prefix", "price_value", "category_value", "brand_value",
          "output_copies", "size_chart_value", "fill_sizes_enabled",
          "minimum_order_quantity_value", "shipping_insurance_value",
          "download_images_enabled"]:
    print(f"  {k} = {settings.get(k)!r}")

print()
products = read_source(src)
print(f"--- READ SOURCE ---")
print(f"  products: {len(products)}")
print(f"  total variants: {sum(len(p.variants) for p in products)}")
for p in products[:3]:
    print(f"  P: {p.product_name[:60]}")
    print(f"     category={p.category!r} brand={p.brand!r}")
    print(f"     variants={len(p.variants)} images={len(p.images)}")
    print(f"     var1[0..2]: {[v.var1 for v in p.variants[:3]]}")

print()
template = Path("/Users/apple/WorkBuddy/2026-09-08-17-31-48/tk-id-converter-repo/assets/batch-product-source.xlsx")

logs = []
def log(s): logs.append(s)
def progress(p, msg): pass

result = convert_source(
    source_xlsx=src,
    output_dir=out_dir,
    settings=settings,
    template_src=template,
    progress=progress,
    log=log,
    download_imgs=False,  # 跳过下载
)
print(f"--- CONVERT RESULT ---")
print(f"  output_paths: {[p.name for p in result.output_paths]}")
print(f"  products: {result.product_count}, rows: {result.row_count}")

print()
print("--- INSPECT OUTPUT ---")
out_xlsx = result.output_paths[0]
import openpyxl
wb = openpyxl.load_workbook(out_xlsx)
ws = wb["Template"]
hdrs = [str(c.value or "") for c in ws[1]]
print(f"  Template dims: {ws.max_row} rows × {ws.max_column} cols")
print(f"  column count: {len(hdrs)}")
print(f"  has pre_order_time? {'pre_order_time' in hdrs}")
print(f"  has minimum_order_quantity? {'minimum_order_quantity' in hdrs}")
print(f"  has shipping_insurance? {'shipping_insurance' in hdrs}")
print(f"  headers: {hdrs}")
print()
print("--- DATA ROWS (first 3) ---")
for r in range(2, 5):
    cells = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None and v != "":
            cells.append(f"[{ws.cell(1,c).value}]={str(v)[:35]}")
    print(f"  R{r}: {' | '.join(cells)}")

print()
# Check key fields across all rows
from collections import Counter
cats = Counter()
colors = Counter()
ship_ins = Counter()
moq = Counter()
for r in range(2, ws.max_row + 1):
    cats[ws.cell(r, hdrs.index("category")+1).value] += 1
    colors[ws.cell(r, hdrs.index("property_value_1")+1).value] += 1
    ship_ins[ws.cell(r, hdrs.index("shipping_insurance")+1).value] += 1
    moq[ws.cell(r, hdrs.index("minimum_order_quantity")+1).value] += 1
print(f"  categories: {dict(cats)}")
print(f"  property_value_1 colors: {dict(colors)}")
print(f"  shipping_insurance: {dict(ship_ins)}")
print(f"  minimum_order_quantity: {dict(moq)}")
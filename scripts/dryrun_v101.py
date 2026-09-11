"""ID v1.0.1 dry-run：用真实 EasyBoss ID 源表验证修复后转化。

验证关键点:
1. pre_order_time 列保留（不删除）— 修复印尼后台"20 个产品报错"
2. seller_sku 不再为空 — ID 源表"平台SKU"列常为 None, 自动生成
3. shipping_insurance / minimum_order_quantity 写入正确
"""
from __future__ import annotations
import sys, shutil
from pathlib import Path

sys.path.insert(0, "/Users/apple/WorkBuddy/2026-09-08-17-31-48/tk-id-converter-repo")

from app.source_reader import read_source
from app.converter import convert_source
from app.config import default_config

src = Path("/Users/apple/Desktop/导出#SKU_2026_09_11_17_15_46.xlsx")
out_dir = Path("/tmp/id_dryrun_v101_out")
if out_dir.exists():
    shutil.rmtree(out_dir)

settings = default_config()["product_xlsx_settings"]
print("--- ID DEFAULT SETTINGS ---")
for k in ["title_prefix", "price_value", "category_value", "brand_value",
          "output_copies", "size_chart_value", "fill_sizes_enabled",
          "minimum_order_quantity_value", "shipping_insurance_value",
          "download_images_enabled"]:
    print(f"  {k} = {settings.get(k)!r}")

print()
products = read_source(src)
print(f"--- READ SOURCE ---")
print(f"  products: {len(products)}")
print(f"  total variants: {sum(len(p.variants) for p in products)}")

template = Path("/Users/apple/WorkBuddy/2026-09-08-17-31-48/tk-id-converter-repo/assets/batch-product-source.xlsx")

logs = []
def log(s): logs.append(s)
def progress(p, msg): pass

result = convert_source(
    source_xlsx=src,
    output_dir=out_dir,
    settings=settings,
    template_src=template,
    progress=progress,
    log=log,
    download_imgs=False,
)
print(f"--- CONVERT RESULT ---")
print(f"  output_paths: {[p.name for p in result.output_paths]}")
print(f"  products: {result.product_count}, rows: {result.row_count}")

print()
print("--- INSPECT OUTPUT ---")
out_xlsx = result.output_paths[0]
import openpyxl
wb = openpyxl.load_workbook(out_xlsx)
ws = wb["Template"]
hdrs = [str(c.value or "") for c in ws[1]]
print(f"  Template dims: {ws.max_row} rows × {ws.max_column} cols")
print(f"  column count: {len(hdrs)}")
print(f"  has pre_order_time? {'pre_order_time' in hdrs}  ← 必须为 True (v1.0.1 修复)")
print(f"  has minimum_order_quantity? {'minimum_order_quantity' in hdrs}")
print(f"  has shipping_insurance? {'shipping_insurance' in hdrs}")
print()

print("--- DATA ROWS (first 3) ---")
for r in range(2, 5):
    cells = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None and v != "":
            cells.append(f"[{ws.cell(1,c).value}]={str(v)[:35]}")
    print(f"  R{r}: {' | '.join(cells)}")

print()
from collections import Counter
cats = Counter()
colors = Counter()
ship_ins = Counter()
moq = Counter()
empty_skus = 0
empty_size_chart = 0
for r in range(2, ws.max_row + 1):
    cats[ws.cell(r, hdrs.index("category")+1).value] += 1
    colors[ws.cell(r, hdrs.index("property_value_1")+1).value] += 1
    ship_ins[ws.cell(r, hdrs.index("shipping_insurance")+1).value] += 1
    moq[ws.cell(r, hdrs.index("minimum_order_quantity")+1).value] += 1
    sku = ws.cell(r, hdrs.index("seller_sku")+1).value
    if not sku:
        empty_skus += 1
    sc = ws.cell(r, hdrs.index("size_chart")+1).value
    if not sc:
        empty_size_chart += 1

print(f"  categories: {dict(cats)}")
print(f"  property_value_1 colors: {dict(colors)}")
print(f"  shipping_insurance: {dict(ship_ins)}")
print(f"  minimum_order_quantity: {dict(moq)}")
print(f"  EMPTY seller_sku: {empty_skus} / {ws.max_row - 1}  ← 必须为 0 (v1.0.1 修复)")
print(f"  EMPTY size_chart: {empty_size_chart} / {ws.max_row - 1}  ← 可接受 (后台可不传)")
